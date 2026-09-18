"""Extract plain text from an uploaded knowledge-base file, dispatching on
its extension. This module's only job is "bytes -> text"; the extracted
text then feeds the same LLM extraction pipeline used by pasted-text
ingest (app.services.extraction_agent).
"""
import io

import openpyxl
from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {"txt", "md", "pdf", "docx", "xlsx"}


def extract_text(filename: str, content: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"unsupported file type: '.{ext}'")

    try:
        if ext in ("txt", "md"):
            text = content.decode("utf-8")
        elif ext == "pdf":
            text = _extract_pdf(content)
        elif ext == "docx":
            text = _extract_docx(content)
        else:  # xlsx
            text = _extract_xlsx(content)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"failed to parse '.{ext}' file: {e}") from e

    if not text.strip():
        raise ValueError(f"no extractable text found in '.{ext}' file")
    return text


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(p for p in pages if p.strip())


def _extract_docx(content: bytes) -> str:
    doc = DocxDocument(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n\n".join(parts)


def _extract_xlsx(content: bytes) -> str:
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)
