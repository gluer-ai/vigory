"""Unit tests for extract_text: bytes -> text per file format. PDF
fixtures are hand-built minimal PDFs (no PDF-writing library needed);
docx/xlsx fixtures are built with the same libraries extract_text reads
them with — a realistic round-trip, since real uploads are themselves
OOXML files produced by Word/Excel/python-docx/openpyxl.
"""
import io

import openpyxl
import pytest
from docx import Document as DocxDocument

from app.services.document_parser import extract_text


def _build_minimal_pdf(text: str) -> bytes:
    """Hand-built single-page PDF with one extractable text stream —
    avoids a PDF-writing dependency just for test fixtures."""
    content_stream = f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 200 200] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content_stream)).encode() + b" >>\nstream\n"
        + content_stream + b"\nendstream",
    ]
    header = b"%PDF-1.4\n"
    body = bytearray()
    offsets = [0]
    pos = len(header)
    for i, obj in enumerate(objects, start=1):
        offsets.append(pos)
        piece = f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
        body += piece
        pos += len(piece)
    xref_offset = pos
    xref = f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets[1:]:
        xref += f"{off:010d} 00000 n \n".encode()
    trailer = (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF"
    ).encode()
    return header + bytes(body) + xref + trailer


def _build_docx(paragraphs: list[str]) -> bytes:
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_xlsx(rows: list[list[str]]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Contacts"
    for row in rows:
        sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def test_extract_text_txt():
    assert extract_text("notes.txt", b"hello world") == "hello world"


def test_extract_text_md():
    assert extract_text("notes.md", b"# Heading\n\nBody text") == "# Heading\n\nBody text"


def test_extract_text_pdf():
    content = _build_minimal_pdf("Hello World")
    assert "Hello World" in extract_text("report.pdf", content)


def test_extract_text_docx():
    content = _build_docx(["First paragraph.", "Second paragraph."])
    text = extract_text("report.docx", content)
    assert "First paragraph." in text
    assert "Second paragraph." in text


def test_extract_text_xlsx():
    content = _build_xlsx([["Name", "Role"], ["Ivan Petrov", "Commander"]])
    text = extract_text("roster.xlsx", content)
    assert "# Sheet: Contacts" in text
    assert "Ivan Petrov" in text
    assert "Commander" in text


def test_extract_text_rejects_unsupported_extension():
    with pytest.raises(ValueError, match="unsupported"):
        extract_text("notes.csv", b"a,b,c")


def test_extract_text_rejects_corrupt_pdf():
    with pytest.raises(ValueError, match="failed to parse"):
        extract_text("broken.pdf", b"not a real pdf")


def test_extract_text_rejects_corrupt_docx():
    with pytest.raises(ValueError, match="failed to parse"):
        extract_text("broken.docx", b"not a real docx")


def test_extract_text_rejects_empty_txt():
    with pytest.raises(ValueError, match="no extractable text"):
        extract_text("empty.txt", b"   ")
